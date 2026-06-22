"""Parse user-uploaded files (PDF, DOCX, TXT/MD) into chunked LangChain Documents
so the chatbot can answer over attachments alongside the PIRLS corpus.

Framework-agnostic: takes (filename, bytes) pairs, so the Streamlit UI can pass
`(f.name, f.getvalue())` for each uploaded file. Chunk size mirrors the corpus
(1000/100) so uploaded chunks rerank fairly against indexed ones.
"""
import io
from typing import List, Tuple

from langchain_core.documents import Document

CHUNK_SIZE = 1000
CHUNK_OVERLAP = 100


def _pdf_text(data: bytes) -> List[Tuple[int, str]]:
    """Return [(page_number, text)] for a PDF."""
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    out = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        if text.strip():
            out.append((i, text))
    return out


def _docx_text(data: bytes) -> str:
    """Extract paragraph + table text from a DOCX (python-docx; xml fallback)."""
    try:
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    except Exception:
        import re
        import zipfile
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            xml = z.read("word/document.xml").decode("utf-8", "ignore")
        xml = re.sub(r"</w:p>", "\n", xml)
        return re.sub(r"<[^>]+>", "", xml)


def _splitter():
    from langchain_text_splitters import RecursiveCharacterTextSplitter
    return RecursiveCharacterTextSplitter(chunk_size=CHUNK_SIZE, chunk_overlap=CHUNK_OVERLAP)


def file_to_documents(filename: str, data: bytes) -> List[Document]:
    """Parse + chunk one uploaded file into Documents tagged with its source name."""
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    splitter = _splitter()
    docs: List[Document] = []
    if ext == "pdf":
        for page, text in _pdf_text(data):
            for chunk in splitter.split_text(text):
                docs.append(Document(page_content=chunk,
                                     metadata={"source": filename, "page": page}))
    else:
        if ext == "docx":
            text = _docx_text(data)
        else:  # txt, md, csv, or anything decodable as text
            text = data.decode("utf-8", "ignore")
        for chunk in splitter.split_text(text):
            docs.append(Document(page_content=chunk, metadata={"source": filename}))
    return docs


def files_to_documents(files: List[Tuple[str, bytes]]) -> List[Document]:
    """Parse + chunk many (filename, bytes) pairs into one Document list."""
    out: List[Document] = []
    for name, data in files:
        try:
            out.extend(file_to_documents(name, data))
        except Exception as exc:  # one bad file shouldn't kill the upload
            out.append(Document(page_content=f"[Could not parse {name}: {exc}]",
                                metadata={"source": name, "error": True}))
    return out


SUPPORTED_EXTENSIONS = ["pdf", "docx", "txt", "md", "csv"]
