"""Build reports/revised_final_report_appendix_C.xlsx from the results/ eval CSVs.

A Summary sheet (one row per run: generator, config, the five metrics overall,
the human/synthetic answer-correctness split, latency, n) plus one detail sheet
per run holding the per-row scores. Reduced-column no-retrieval/ablation files
(answer_correctness + gold_context_similarity only) are handled.
"""
import csv, statistics, re
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).parent.parent          # repo root
RESULTS = ROOT / "results"
OUT = Path(__file__).parent / "revised_final_report_appendix_C.xlsx"

METRICS = ["contextual_precision", "contextual_recall", "faithfulness",
           "answer_correctness", "gold_context_similarity"]

# run label -> (eval csv, generator, configuration, latency)
RUNS = [
    ("Basic (baseline)",            "eval_r3_standard_ollama_deepeval.csv",                      "llama3:8b",      "k=4, 1000/100",                 "19.5 s"),
    ("Advanced v1",                 "eval_r3_crag_ollama_deepeval.csv",                          "llama3:8b",      "gemma3:1b grader",              "74.1 s"),
    ("Advanced v2",                 "eval_r3_cragpp_ollama_deepeval.csv",                        "llama3:8b",      "per-sub-q retrieve + dedup",    "101.8 s"),
    ("Basic + rerank",              "eval_r3_standard_ollama_rerank_deepeval.csv",               "llama3:8b",      "k=20 -> rerank -> 4",           "21.9 s"),
    ("chunk-512",                   "eval_r3_standard_ollama_chunk512_deepeval.csv",             "llama3:8b",      "k=4, 512/64",                   "11.3 s"),
    ("Basic, deepseek",             "eval_r3_standard_ollama_deepseek_deepeval.csv",             "deepseek-r1:8b", "k=4, 1000/100",                 "62.7 s"),
    ("Basic + rerank, gpt-5.4",     "eval_r3_standard_openai-mini_rerank_deepeval.csv",          "gpt-5.4-mini",   "k=20 -> rerank -> 4",           "1.6 s*"),
    ("rerank + extract prompt",     "eval_r3_standard_ollama_rerank_extract_deepeval.csv",       "llama3:8b",      "extract-style prompt",          "19.4 s"),
    ("dense + bge reranker",        "eval_r3_standard_ollama_rerank_bgebase_deepeval.csv",       "llama3:8b",      "k=20 -> bge-base -> 4",         "~22 s"),
    ("hybrid + bge (best local)",   "eval_r3_standard_ollama_hybrid_bgebase_deepeval.csv",       "llama3:8b",      "BM25+dense -> RRF -> bge -> 4", "~23 s"),
    ("hybrid + MiniLM",             "eval_r3_standard_ollama_hybrid_minilm_deepeval.csv",        "llama3:8b",      "BM25+dense -> RRF -> MiniLM -> 4", "19.2 s"),
    ("hybrid + bge, deepseek",      "eval_r3_standard_ollama_hybrid_deepseek_bgebase_deepeval.csv","deepseek-r1:8b","BM25+dense -> RRF -> bge -> 4", "62.4 s"),
    ("hybrid + bge, gpt-5.4 (best overall)", "eval_r3_standard_openai-mini_hybrid_bgebase_deepeval.csv", "gpt-5.4-mini", "BM25+dense -> RRF -> bge -> 4", "7.9 s*"),
    ("Advanced v3, llama3:8b",      "eval_r3_cragpp_ollama_hybrid_cragfix_deepeval.csv",         "llama3:8b",      "hybrid + 4 fixes (incomplete)", "-"),
    ("Advanced v3, gpt-5.4-mini",   "eval_r3_cragpp_openai-mini_hybrid_cragfix_deepeval.csv",     "gpt-5.4-mini",   "hybrid + 4 fixes",              "9.9 s*"),
    ("Closed-book, llama3:8b",      "eval_r3_noretrieval_ollama_deepeval.csv",                   "llama3:8b",      "no retrieval",                  "-"),
    ("Closed-book, gpt-5.4-mini",   "eval_r3_noretrieval_openai-mini_deepeval.csv",              "gpt-5.4-mini",   "no retrieval",                  "-"),
    ("Retrieval-only",              "eval_r3_retrievalonly_hybrid_bgebase_deepeval.csv",         "(none)",         "top-4 hybrid chunks as answer", "-"),
]

def mean(rows, key):
    vals = [float(r[key]) for r in rows if r.get(key) not in (None, "", "nan")]
    return round(statistics.mean(vals), 3) if vals else None

def safe_sheet(name):
    name = re.sub(r'[:\\/?*\[\]]', "-", name)
    return name[:31]

HEAD_FILL = PatternFill("solid", fgColor="0F4761")
HEAD_FONT = Font(bold=True, color="FFFFFF")

wb = openpyxl.Workbook()
summary = wb.active
summary.title = "Summary"
hdr = ["Run", "Generator", "Configuration", "n",
       "Ctx Precision", "Ctx Recall", "Faithfulness", "Answer Correctness",
       "Gold-Ctx Sim", "AC (human)", "AC (synthetic)", "Latency"]
summary.append(hdr)
for c in summary[1]:
    c.fill = HEAD_FILL; c.font = HEAD_FONT; c.alignment = Alignment(horizontal="center")

for label, fname, gen, config, latency in RUNS:
    fp = RESULTS / fname
    if not fp.exists():
        print("MISSING", fname); continue
    rows = list(csv.DictReader(open(fp, encoding="utf-8")))
    human = [r for r in rows if r.get("data_type") == "human"]
    synth = [r for r in rows if r.get("data_type") == "synthetic"]
    summary.append([
        label, gen, config, len(rows),
        mean(rows, "contextual_precision"), mean(rows, "contextual_recall"),
        mean(rows, "faithfulness"), mean(rows, "answer_correctness"),
        mean(rows, "gold_context_similarity"),
        mean(human, "answer_correctness"), mean(synth, "answer_correctness"),
        latency,
    ])
    # detail sheet
    ws = wb.create_sheet(safe_sheet(label))
    cols = [c for c in ["row_id", "system", "generator", "data_type"] + METRICS
            if c in (rows[0].keys() if rows else [])]
    ws.append(cols)
    for c in ws[1]:
        c.fill = HEAD_FILL; c.font = HEAD_FONT
    for r in rows:
        out = []
        for c in cols:
            v = r.get(c, "")
            if c in METRICS and v not in (None, "", "nan"):
                v = float(v)
            out.append(v)
        ws.append(out)
    ws.freeze_panes = "A2"

# column widths + freeze on summary
for col, w in zip("ABCDEFGHIJKL", [34, 15, 30, 6, 13, 12, 13, 17, 13, 12, 14, 10]):
    summary.column_dimensions[col].width = w
summary.freeze_panes = "A2"

wb.save(OUT)
print("wrote", OUT, "| sheets:", len(wb.sheetnames))
